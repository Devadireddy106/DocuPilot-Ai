"""Excel workbook export.

Sheets: Extracted_Data (+ Extracted_Data_2.. for further independent tables), Validation_Report,
Processing_Summary. The workbook is written to a temp file, re-opened and verified, and only then
moved to its final location. Untrusted strings can never become formulas.
"""
from __future__ import annotations

import logging
import re
import tempfile
import zipfile
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config.settings import SETTINGS, Settings
from src import file_manager
from src.models import ExtractionResult, Issue, Table
from src.processing_report import summary_rows, table_index_rows, page_rows
from src.schema_normalizer import infer_column_types, to_typed_value

log = logging.getLogger(__name__)

NAVY, BLUE, EMERALD, AMBER, RED, LIGHT = "0B1F3A", "1D4ED8", "059669", "FEF3C7", "FEE2E2", "F1F5F9"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
FLAG_FILL = PatternFill("solid", fgColor=AMBER)
ERR_FILL = PatternFill("solid", fgColor=RED)
THIN = Side(style="thin", color="CBD5E1")

_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r", "\n")
_NUMERIC_LIKE = re.compile(r"^[+-]?(\d+(\.\d*)?|\.\d+)([eE][+-]?\d+)?$")


def safe_text(value: Any, max_chars: int = SETTINGS.max_cell_chars) -> str:
    """Return a string that is safe to store in a worksheet cell (never evaluated as a formula)."""
    s = "" if value is None else str(value)
    s = _ILLEGAL_XML.sub("", s)
    return s[:max_chars]


def _needs_quote_prefix(s: str) -> bool:
    return bool(s) and s[0] in _FORMULA_START


def put_text(ws, row: int, col: int, value: Any, max_chars: int = SETTINGS.max_cell_chars):
    """Write text defensively: stored as a string cell with Excel's 'quote prefix' flag when it starts
    with a formula-trigger character, so Excel/LibreOffice show it as text and never evaluate it."""
    s = safe_text(value, max_chars)
    cell = ws.cell(row=row, column=col)
    cell.value = s
    cell.data_type = "s"  # force string even for '=...'
    if _needs_quote_prefix(s):
        cell.quotePrefix = True
    return cell


def _style_header(ws, ncols: int, row: int = 1) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)
    ws.row_dimensions[row].height = 24


def _autowidth(ws, ncols: int, min_w: int = 9, max_w: int = 48, sample_rows: int = 200) -> None:
    for c in range(1, ncols + 1):
        longest = 0
        for r in range(1, min(ws.max_row, sample_rows) + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None:
                longest = max(longest, len(str(v)))
        ws.column_dimensions[get_column_letter(c)].width = max(min_w, min(max_w, longest + 2))


def _issues_by_row(issues: List[Issue]) -> Dict[Tuple[str, int], List[Issue]]:
    out: Dict[Tuple[str, int], List[Issue]] = defaultdict(list)
    for i in issues:
        if i.table_id and i.row:
            out[(i.table_id, i.row)].append(i)
    return out


def _unique(name: str, taken: List[str]) -> str:
    cand, k = name, 2
    low = {t.lower() for t in taken}
    while cand.lower() in low:
        cand, k = f"{name} ({k})", k + 1
    return cand


def _write_table_sheet(ws, t: Table, issues: List[Issue]) -> None:
    types = infer_column_types(t)
    meta_names: List[str] = []
    for want in ("Source_Page", "Extraction_Method", "Review_Flag"):
        meta_names.append(_unique(want, t.headers + meta_names))
    all_headers = list(t.headers) + meta_names
    for c, h in enumerate(all_headers, start=1):
        put_text(ws, 1, c, h, 255)
    _style_header(ws, len(all_headers))

    by_row = _issues_by_row(issues)
    n = len(t.headers)
    for ri, row in enumerate(t.rows):
        r = ri + 2
        for ci in range(n):
            raw = row[ci] if ci < len(row) else ""
            val, fmt = to_typed_value(raw, types[ci])
            if val is None:
                continue
            if isinstance(val, str):
                put_text(ws, r, ci + 1, val)
            else:
                cell = ws.cell(row=r, column=ci + 1, value=val)
                if fmt:
                    cell.number_format = fmt
                if types[ci].kind == "numeric":
                    cell.alignment = Alignment(horizontal="right")
        ws.cell(row=r, column=n + 1, value=t.row_pages[ri])
        put_text(ws, r, n + 2, t.row_methods[ri])
        found = by_row.get((t.table_id, ri + 1), [])
        flagged = [i for i in found if i.severity in ("Error", "Warning")]
        if flagged:
            put_text(ws, r, n + 3, "Review: " + ", ".join(i.issue_id for i in flagged))
            fill = ERR_FILL if any(i.severity == "Error" for i in flagged) else FLAG_FILL
            for ci in range(1, n + 4):
                ws.cell(row=r, column=ci).fill = fill
        elif found:  # Info only (e.g. auto-corrected)
            put_text(ws, r, n + 3, "Note: " + ", ".join(i.issue_id for i in found))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(all_headers))}{max(2, len(t.rows) + 1)}"
    _autowidth(ws, len(all_headers))


def _write_empty_sheet(ws, result: ExtractionResult) -> None:
    put_text(ws, 1, 1, "Status")
    put_text(ws, 1, 2, "Message")
    _style_header(ws, 2)
    put_text(ws, 2, 1, result.status)
    put_text(ws, 2, 2, "No tables were extracted from this document. " + (result.status_detail or "") + " See Validation_Report and Processing_Summary.")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 18, 100
    ws.freeze_panes = "A2"


def _write_validation(ws, result: ExtractionResult) -> None:
    heads = ["Issue ID", "Severity", "Category", "Page", "Sheet", "Row Reference", "Column Name",
             "Issue Description", "Suggested Action", "Resolution Status"]
    for c, h in enumerate(heads, start=1):
        put_text(ws, 1, c, h)
    _style_header(ws, len(heads))
    for r, i in enumerate(result.issues, start=2):
        put_text(ws, r, 1, i.issue_id)
        put_text(ws, r, 2, i.severity)
        put_text(ws, r, 3, i.category)
        if i.page:
            ws.cell(row=r, column=4, value=i.page)
        put_text(ws, r, 5, i.sheet or "")
        put_text(ws, r, 6, f"Row {i.excel_row}" if i.excel_row else "")
        put_text(ws, r, 7, i.column or "")
        put_text(ws, r, 8, i.description)
        put_text(ws, r, 9, i.action)
        put_text(ws, r, 10, i.status)
        fill = ERR_FILL if i.severity == "Error" else FLAG_FILL if i.severity == "Warning" else None
        if fill:
            ws.cell(row=r, column=2).fill = fill
        ws.cell(row=r, column=8).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=9).alignment = Alignment(wrap_text=True, vertical="top")
    if not result.issues:
        put_text(ws, 2, 1, "-")
        put_text(ws, 2, 8, "No validation findings. This does not guarantee the extraction is correct; automated checks are limited.")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{max(2, len(result.issues) + 1)}"
    for col, w in zip("ABCDEFGHIJ", (9, 10, 12, 7, 18, 14, 18, 70, 48, 16)):
        ws.column_dimensions[col].width = w


def _write_summary(ws, result: ExtractionResult) -> None:
    r = 1
    if result.status != "Complete":
        banner = {
            "Partial": "PARTIAL RESULT - some pages could not be processed. " + result.status_detail,
            "Failed": "FAILED - " + result.status_detail,
            "No tables found": "NO TABLES FOUND - " + result.status_detail,
            "Complete (review needed)": "REVIEW NEEDED - " + result.status_detail,
        }.get(result.status, result.status_detail)
        put_text(ws, r, 1, banner).font = Font(bold=True, color="991B1B")
        ws.cell(row=r, column=1).fill = ERR_FILL
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 36
        r += 2
    put_text(ws, r, 1, "Metric"); put_text(ws, r, 2, "Value")
    _style_header(ws, 2, r)
    for k, v in summary_rows(result):
        r += 1
        put_text(ws, r, 1, k)
        if isinstance(v, (int, float)):
            ws.cell(row=r, column=2, value=v)
        else:
            put_text(ws, r, 2, v)
        ws.cell(row=r, column=2).alignment = Alignment(horizontal="left", wrap_text=True, vertical="top")

    idx = table_index_rows(result)
    if idx:
        r += 2
        for c, h in enumerate(["Sheet", "Table", "Pages", "Rows", "Columns"], start=1):
            put_text(ws, r, c, h)
        _style_header(ws, 5, r)
        for row in idx:
            r += 1
            for c, v in enumerate(row, start=1):
                put_text(ws, r, c, v) if isinstance(v, str) else ws.cell(row=r, column=c, value=v)

    r += 2
    for c, h in enumerate(["Page", "Page type", "Status", "Method kept", "Attempts", "Notes"], start=1):
        put_text(ws, r, c, h)
    _style_header(ws, 6, r)
    for row in page_rows(result):
        r += 1
        for c, v in enumerate(row, start=1):
            put_text(ws, r, c, v) if isinstance(v, str) else ws.cell(row=r, column=c, value=v)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 40
    for col, w in zip("CDEF", (16, 22, 40, 70)):
        ws.column_dimensions[col].width = w


def build_workbook(result: ExtractionResult) -> Workbook:
    wb = Workbook()
    ws0 = wb.active
    if result.tables:
        for k, t in enumerate(result.tables):
            ws = ws0 if k == 0 else wb.create_sheet()
            ws.title = t.sheet_name or ("Extracted_Data" if k == 0 else f"Extracted_Data_{k + 1}")
            _write_table_sheet(ws, t, result.issues)
    else:
        ws0.title = "Extracted_Data"
        _write_empty_sheet(ws0, result)
    _write_validation(wb.create_sheet("Validation_Report"), result)
    _write_summary(wb.create_sheet("Processing_Summary"), result)
    wb.properties.title = "DocuPilot AI extraction"
    wb.properties.creator = "DocuPilot AI"
    wb.properties.description = f"Extracted from {file_manager.display_name(result.filename)} - status: {result.status}"
    return wb


def verify_workbook(path: Path, result: ExtractionResult) -> List[str]:
    """Re-open the saved file and check structure. Returns a list of problems (empty = OK)."""
    problems: List[str] = []
    try:
        with zipfile.ZipFile(path) as z:
            bad = z.testzip()
            if bad:
                problems.append(f"corrupt zip member {bad}")
        wb = load_workbook(path)  # full (not read-only) load so formulas/types are inspectable
    except Exception as exc:
        return [f"workbook cannot be re-opened: {type(exc).__name__}: {exc}"]
    names = wb.sheetnames
    for need in ("Extracted_Data", "Validation_Report", "Processing_Summary"):
        if need not in names:
            problems.append(f"missing sheet {need}")
    if len({n.lower() for n in names}) != len(names):
        problems.append("duplicate sheet names")
    for t in result.tables:
        if t.sheet_name in names:
            ws = wb[t.sheet_name]
            if ws.max_row - 1 != t.n_rows:
                problems.append(f"{t.sheet_name}: expected {t.n_rows} data rows, found {ws.max_row - 1}")
        else:
            problems.append(f"missing table sheet {t.sheet_name}")
    for ws in wb.worksheets:  # no cell may be a live formula
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == "f":
                    problems.append(f"{ws.title}!{c.coordinate} is a formula")
                    break
    wb.close()
    return problems


def export_workbook(result: ExtractionResult, out_dir: Optional[Path] = None, settings: Settings = SETTINGS) -> Path:
    """Write, verify and publish the workbook. Raises RuntimeError if verification fails."""
    out_dir = Path(out_dir or settings.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    wb = build_workbook(result)
    stem = file_manager.sanitize_filename(result.filename)
    with tempfile.TemporaryDirectory(prefix="docupilot_xlsx_") as td:
        tmp = Path(td) / "workbook.xlsx"
        wb.save(tmp)
        problems = verify_workbook(tmp, result)
        if problems:
            raise RuntimeError("The generated workbook failed verification: " + "; ".join(problems[:5]))
        from datetime import datetime
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return file_manager.publish_file(tmp, out_dir, f"{file_manager.OUTPUT_PREFIX}{stem}_{stamp}", ".xlsx")

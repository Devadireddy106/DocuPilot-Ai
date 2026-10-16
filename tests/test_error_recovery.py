import dataclasses
import json
import tempfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

import fixture_factory as ff
from src import file_manager
from src.extraction_agent import ExtractionAgent
from src.pipeline import convert_pdf


def _work_dirs():
    return set(Path(tempfile.gettempdir()).glob("docupilot_work_*"))


def test_one_failing_page_does_not_stop_the_document(invoice_pdf, settings, monkeypatch):
    original = ExtractionAgent._execute_plan

    def flaky(self, page, pr, ctx):
        if pr.page == 2:
            raise RuntimeError("simulated crash")
        return original(self, page, pr, ctx)

    monkeypatch.setattr(ExtractionAgent, "_execute_plan", flaky)
    r = ExtractionAgent(settings).run(invoice_pdf)
    assert [p.status for p in r.page_results] == ["ok", "failed"]
    assert "simulated crash" in r.page_results[1].failure_reason
    assert r.status == "Partial" and r.tables and r.tables[0].pages == [1]
    assert any(i.check == "failed_page" and i.page == 2 for i in r.issues)


def test_strategy_error_triggers_fallback_strategy(invoice_unruled_pdf, settings, monkeypatch):
    from src import pdf_extractor

    def boom(*a, **k):
        raise RuntimeError("lines strategy exploded")

    monkeypatch.setattr(pdf_extractor, "extract_pdfplumber_lines", boom)
    r = ExtractionAgent(settings).run(invoice_unruled_pdf)
    p = r.page_results[0]
    assert p.status == "ok" and p.method != "pdfplumber_lines"
    assert p.attempts[0].success is False and "exploded" in p.attempts[0].reason


def test_time_budget_exhaustion_marks_remaining_pages_failed(invoice_pdf, settings):
    s = dataclasses.replace(settings, max_processing_seconds=0)
    r = ExtractionAgent(s).run(invoice_pdf)
    assert r.aborted_for_time
    assert all(p.status == "failed" for p in r.page_results)
    assert r.status == "Failed"
    assert any(i.check == "time_budget" for i in r.issues)


def test_progress_callback_errors_never_break_processing(small_digital_pdf, settings):
    def bad(frac, msg):
        raise ValueError("ui exploded")

    r = ExtractionAgent(settings, bad).run(small_digital_pdf)
    assert r.status == "Complete"


def test_progress_is_monotonic_and_ends_at_one(small_digital_pdf, settings):
    seen = []
    out = convert_pdf(small_digital_pdf, "a.pdf", progress=lambda f, m: seen.append(f), settings=settings, out_dir=settings.output_dir)
    assert out.ok and seen == sorted(seen) and seen[-1] == 1.0


@pytest.mark.parametrize("maker,fragment", [
    (ff.make_corrupt_pdf, "PDF"),
    (ff.make_not_pdf, "PDF"),
])
def test_bad_inputs_get_friendly_messages(maker, fragment, pdf_dir, settings):
    p = maker(pdf_dir / f"bad_{maker.__name__}.pdf")
    out = convert_pdf(p, p.name, settings=settings, out_dir=settings.output_dir)
    assert not out.ok and fragment in out.message and "Traceback" not in out.message
    assert list(settings.output_dir.iterdir()) == []     # nothing half-written


def test_encrypted_and_oversized_inputs_via_pipeline(pdf_dir, small_digital_pdf, settings):
    enc = ff.make_encrypted_pdf(small_digital_pdf, pdf_dir / "enc2.pdf")
    assert "password" in convert_pdf(enc, "enc2.pdf", settings=settings, out_dir=settings.output_dir).message.lower()
    big = ff.make_many_pages_pdf(pdf_dir / "p51b.pdf", 51)
    assert "50" in convert_pdf(big, "p51b.pdf", settings=settings, out_dir=settings.output_dir).message


def test_empty_extraction_returns_valid_workbook_and_says_so(prose_pdf, settings):
    out = convert_pdf(prose_pdf, "notes.pdf", settings=settings, out_dir=settings.output_dir)
    assert out.ok and out.result.status == "No tables found"
    wb = load_workbook(out.workbook_path)
    assert "No tables were extracted" in wb["Extracted_Data"]["B2"].value


def test_temp_workspace_cleaned_after_success_and_failure(small_digital_pdf, pdf_dir, settings):
    before = _work_dirs()
    convert_pdf(small_digital_pdf, "ok.pdf", settings=settings, out_dir=settings.output_dir)
    convert_pdf(ff.make_corrupt_pdf(pdf_dir / "c2.pdf"), "c2.pdf", settings=settings, out_dir=settings.output_dir)
    assert _work_dirs() == before
    assert not list(Path(tempfile.gettempdir()).glob("docupilot_xlsx_*"))


def test_workspace_removed_even_if_body_raises():
    with pytest.raises(RuntimeError):
        with file_manager.workspace() as w:
            kept = w
            (w / "x.txt").write_text("hi")
            raise RuntimeError("fail")
    assert not kept.exists()


def test_repeated_runs_are_deterministic_and_independent(invoice_pdf, settings):
    a = ExtractionAgent(settings).run(invoice_pdf)
    b = ExtractionAgent(settings).run(invoice_pdf)
    assert a.tables[0].rows == b.tables[0].rows and a.tables[0].headers == b.tables[0].headers
    assert [i.check for i in a.issues] == [i.check for i in b.issues]
    o1 = convert_pdf(invoice_pdf, "r.pdf", settings=settings, out_dir=settings.output_dir)
    o2 = convert_pdf(invoice_pdf, "r.pdf", settings=settings, out_dir=settings.output_dir)
    assert o1.workbook_path != o2.workbook_path and o1.workbook_path.exists() and o2.workbook_path.exists()


def test_processing_log_json_is_valid_and_private(small_digital_pdf, settings):
    out = convert_pdf(small_digital_pdf, "priv.pdf", settings=settings, out_dir=settings.output_dir)
    data = json.loads(out.log_path.read_text(encoding="utf-8"))
    assert data["status"] == "Complete" and data["pages"][0]["attempts"]
    assert "INV-" not in json.dumps(data)               # cell values never logged


def test_validation_csv_is_formula_safe_and_has_bom(small_digital_pdf, settings):
    from src.processing_report import csv_safe
    assert csv_safe("=1+1").startswith("'") and csv_safe("ok") == "ok"
    out = convert_pdf(small_digital_pdf, "v.pdf", settings=settings, out_dir=settings.output_dir)
    raw = out.validation_csv_path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf") and b"Issue ID" in raw


def test_old_outputs_purged_but_foreign_files_kept(tmp_path):
    import os, time
    old = tmp_path / "docupilot_old.xlsx"; old.write_bytes(b"x")
    foreign = tmp_path / "mine.xlsx"; foreign.write_bytes(b"x")
    past = time.time() - 72 * 3600
    os.utime(old, (past, past)); os.utime(foreign, (past, past))
    assert file_manager.purge_old_outputs(tmp_path, 24) == 1
    assert not old.exists() and foreign.exists()


def test_sanitize_filename():
    assert file_manager.sanitize_filename("../../a b/c?.pdf") == "c"
    assert file_manager.sanitize_filename("CON.pdf") == "_CON"
    assert file_manager.sanitize_filename("") == "document"
    assert file_manager.sanitize_filename("résumé tabelle.pdf") == "resume_tabelle"
    assert file_manager.display_name("a/b/c\x00.pdf") == "c.pdf"

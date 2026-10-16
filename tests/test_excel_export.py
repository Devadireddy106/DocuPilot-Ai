import datetime as dt
import tempfile
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conftest import make_table
from src.excel_exporter import build_workbook, export_workbook, put_text, safe_text, verify_workbook
from src.models import ExtractionResult, PageResult
from src.validator import validate


def result_for(tables, pages=None, status="Complete", filename="report.pdf"):
    pages = pages or [PageResult(page=1, status="ok", method="pdfplumber_lines", kind="digital")]
    for i, t in enumerate(tables, start=1):
        t.table_id = f"T{i}"
        t.sheet_name = "Extracted_Data" if i == 1 else f"Extracted_Data_{i}"
    r = ExtractionResult(filename=filename, sha256="0" * 64, total_pages=len(pages), page_results=pages,
                         tables=list(tables), status=status, status_detail="detail", doc_kind="digital")
    r.issues = validate(tables, pages)
    return r


def sample_table():
    rows = [["001", "2026-03-14", "Widget", "1,234.50", "12.5%"], ["002", "2026-03-15", "Gadget", "99.00", "7.5%"],
            ["003", "2026-03-16", "Gizmo", "5.25", "10.0%"]]
    return make_table(["Code", "Date", "Name", "Amount", "Rate"], rows)


def test_workbook_has_required_sheets_and_is_valid(tmp_path, settings):
    r = result_for([sample_table()])
    path = export_workbook(r, tmp_path)
    assert path.exists() and path.suffix == ".xlsx" and zipfile.is_zipfile(path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Extracted_Data", "Validation_Report", "Processing_Summary"]
    assert verify_workbook(path, r) == []


def test_extracted_data_formatting_types_and_flags(tmp_path):
    r = result_for([sample_table()])
    ws = load_workbook(export_workbook(r, tmp_path))["Extracted_Data"]
    headers = [c.value for c in ws[1]]
    assert headers == ["Code", "Date", "Name", "Amount", "Rate", "Source_Page", "Extraction_Method", "Review_Flag"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref and ws.auto_filter.ref.startswith("A1:")
    assert ws["A2"].value == "001" and isinstance(ws["A2"].value, str)         # leading zeros kept
    assert isinstance(ws["B2"].value, dt.datetime) and ws["B2"].number_format == "yyyy-mm-dd"
    assert ws["D2"].value == pytest.approx(1234.5) and ws["D2"].number_format == "#,##0.00"
    assert ws["E2"].value == pytest.approx(0.125) and ws["E2"].number_format.endswith("%")
    assert ws["F2"].value == 1
    assert ws["A1"].font.bold and ws.column_dimensions["C"].width >= 9


def test_formula_injection_is_neutralised(tmp_path):
    evil = ["=HYPERLINK(\"http://evil\",\"x\")", "+SUM(A1:A9)", "-2+3", "@cmd", "=1+1"]
    rows = [[v, "x"] for v in evil]
    t = make_table(["Payload", "Other"], rows)
    path = export_workbook(result_for([t], filename="=evil.pdf"), tmp_path)
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f", f"{ws.title}!{c.coordinate} became a formula"
    ws = wb["Extracted_Data"]
    for i, v in enumerate(evil, start=2):
        assert ws.cell(row=i, column=1).value == v           # text preserved verbatim
        assert ws.cell(row=i, column=1).quotePrefix is True  # and flagged as text for Excel
    # also inspect raw XML: no <f> formula elements anywhere
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if name.startswith("xl/worksheets/"):
                assert b"<f>" not in z.read(name) and b"<f " not in z.read(name)


def test_filename_in_summary_is_sanitised_and_output_name_safe(tmp_path):
    path = export_workbook(result_for([sample_table()], filename="../../etc/pass wd\x00.pdf"), tmp_path)
    assert path.parent == tmp_path and ".." not in path.name and "\x00" not in path.name
    assert path.name.startswith("docupilot_")


def test_illegal_characters_and_long_text_are_handled(tmp_path):
    t = make_table(["A", "B"], [["bad\x00\x07char", "x" * 40000]])
    path = export_workbook(result_for([t]), tmp_path)
    ws = load_workbook(path)["Extracted_Data"]
    assert ws["A2"].value == "badchar" and len(ws["B2"].value) <= 32000
    assert safe_text(None) == ""


def test_multiple_tables_get_unique_valid_sheet_names(tmp_path):
    ts = [sample_table(), sample_table(), sample_table()]
    r = result_for(ts)
    wb = load_workbook(export_workbook(r, tmp_path))
    assert wb.sheetnames[:3] == ["Extracted_Data", "Extracted_Data_2", "Extracted_Data_3"]
    assert len({n.lower() for n in wb.sheetnames}) == len(wb.sheetnames)
    assert all(len(n) <= 31 for n in wb.sheetnames)


def test_no_tables_still_produces_valid_workbook_with_message(tmp_path):
    r = result_for([], pages=[PageResult(page=1, status="no_tables", kind="digital")], status="No tables found")
    path = export_workbook(r, tmp_path)
    wb = load_workbook(path)
    assert "No tables were extracted" in wb["Extracted_Data"]["B2"].value
    assert wb["Processing_Summary"]["A1"].value.startswith("NO TABLES FOUND")


def test_partial_result_is_marked_partial_everywhere(tmp_path):
    pages = [PageResult(page=1, status="ok", method="m", kind="digital"), PageResult(page=2, status="failed", failure_reason="OCR missing", kind="scanned")]
    r = result_for([sample_table()], pages=pages, status="Partial")
    r.status_detail = "1 page(s) could not be processed"
    wb = load_workbook(export_workbook(r, tmp_path))
    assert wb["Processing_Summary"]["A1"].value.startswith("PARTIAL RESULT")
    vals = {row[0].value: row[1].value for row in wb["Processing_Summary"].iter_rows(min_row=3, max_col=2) if row[0].value}
    assert vals["Completion status"] == "Partial" and vals["Failed pages"] == 1
    assert any(row[1].value == "Error" for row in wb["Validation_Report"].iter_rows(min_row=2))


def test_validation_report_columns_and_summary_fields(tmp_path):
    t = make_table(["Item", "Qty"], [["A", "1"], ["A", "1"]])
    r = result_for([t])
    wb = load_workbook(export_workbook(r, tmp_path))
    heads = [c.value for c in wb["Validation_Report"][1]]
    for need in ("Issue ID", "Severity", "Page", "Row Reference", "Column Name", "Issue Description", "Suggested Action", "Resolution Status"):
        assert need in heads
    labels = {row[0].value for row in wb["Processing_Summary"].iter_rows(max_col=1)}
    for need in ("Input filename", "Total pages", "Successfully processed pages", "Failed pages", "OCR pages (OCR output used)",
                 "Extracted tables", "Extracted records (rows)", "Validation issues (total)", "Processing duration (seconds)",
                 "Extraction methods used", "Completion status"):
        assert need in labels


def test_flagged_rows_reference_issue_ids(tmp_path):
    t = make_table(["Item", "Qty"], [["A", "1"], ["A", "1"]])
    wb = load_workbook(export_workbook(result_for([t]), tmp_path))
    ws = wb["Extracted_Data"]
    assert ws["E3"].value and ws["E3"].value.startswith("Review: V")   # Review_Flag column


def test_verify_workbook_detects_problems(tmp_path):
    r = result_for([sample_table()])
    good = export_workbook(r, tmp_path)
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"not an xlsx")
    assert verify_workbook(bad, r)
    r2 = result_for([sample_table(), sample_table()])           # expects 2 table sheets, file has 1
    assert any("missing table sheet" in p for p in verify_workbook(good, r2))


def test_export_never_overwrites_and_leaves_no_temp_files(tmp_path):
    before = set(Path(tempfile.gettempdir()).glob("docupilot_xlsx_*"))
    r = result_for([sample_table()])
    marker = tmp_path / "unrelated.xlsx"
    marker.write_bytes(b"keep me")
    p1, p2 = export_workbook(r, tmp_path), export_workbook(r, tmp_path)
    assert p1 != p2 or p1.exists()
    assert marker.read_bytes() == b"keep me"
    assert set(Path(tempfile.gettempdir()).glob("docupilot_xlsx_*")) == before
    assert sorted(x.suffix for x in tmp_path.iterdir() if x.name != "unrelated.xlsx") == [".xlsx"] * len(list(tmp_path.glob("docupilot_*.xlsx")))


def test_put_text_returns_text_cell_for_equals_sign():
    from openpyxl import Workbook
    ws = Workbook().active
    c = put_text(ws, 1, 1, "=2+2")
    assert c.data_type == "s" and c.quotePrefix
